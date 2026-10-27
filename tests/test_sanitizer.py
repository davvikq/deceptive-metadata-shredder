from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET

import pytest
from docx import Document
from lxml import etree
from pypdf import PdfReader, PdfWriter
from pypdf.annotations import Text
from pypdf.generic import NameObject, TextStringObject
from dms.core.analyzer import analyze
from dms.core.sanitizer import _safe_extractall, _strip_docx_tracked_changes, remove_all


FIXTURES = Path(__file__).parent / "fixtures"


def test_cleaned_copy_is_created_without_touching_original(tmp_path) -> None:
    source = FIXTURES / "test_no_meta.png"
    output = tmp_path / "cleaned.png"

    result = remove_all(source, output)

    assert result.exists()
    assert source.exists()
    assert source.read_bytes() == (FIXTURES / "test_no_meta.png").read_bytes()
    remaining = [
        field
        for field in analyze(result).fields
        if field.key != "SourceFile" and not field.exiftool_tag.startswith("System:")
    ]

    assert not any(field.is_sensitive for field in remaining)


def test_pdf_cleanup_removes_info_xmp_and_previous_author_bytes(tmp_path) -> None:
    source = tmp_path / "authored.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.add_outline_item("Chapter", 0)
    writer.add_attachment(filename="note.txt", data=b"keep this attachment")
    annotation = Text(rect=(20, 20, 40, 40), text="Keep this note")
    annotation[NameObject("/T")] = TextStringObject("Private Nick")
    writer.add_annotation(page_number=0, annotation=annotation)
    writer.add_metadata({"/Author": "Private Nick", "/Creator": "Private Nick"})
    writer.xmp_metadata = b"<metadata>Private Nick</metadata>"
    writer.write(source)

    cleaned = remove_all(source)
    report = analyze(cleaned)
    reader = PdfReader(cleaned)

    assert source.read_bytes() != cleaned.read_bytes()
    assert b"Private Nick" not in cleaned.read_bytes()
    assert reader.metadata is None
    assert reader.xmp_metadata is None
    assert len(reader.pages) == 1
    assert len(reader.outline) == 1
    assert reader.attachments["note.txt"] == [b"keep this attachment"]
    saved_annotation = reader.pages[0]["/Annots"][0].get_object()
    assert "/T" not in saved_annotation
    assert saved_annotation["/Contents"] == "Keep this note"
    assert not {"Author", "Creator"} & {field.key for field in report.fields}


def test_pdf_fixture_author_is_absent_after_reanalysis(tmp_path) -> None:
    cleaned = remove_all(FIXTURES / "test_document.pdf", tmp_path / "document_cleaned.pdf")

    assert b"Jane Doe" not in cleaned.read_bytes()
    assert not {"Author", "Creator", "Producer"} & {field.key for field in analyze(cleaned).fields}


def test_docx_cleanup_removes_author_from_properties_and_comments(tmp_path) -> None:
    source = tmp_path / "authored.docx"
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    content_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    app_ns = "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
    with zipfile.ZipFile(FIXTURES / "test_document.docx") as original, zipfile.ZipFile(source, "w") as authored:
        for info in original.infolist():
            payload = original.read(info.filename)
            if info.filename in {"_rels/.rels", "word/_rels/document.xml.rels", "[Content_Types].xml", "docProps/app.xml"}:
                root = etree.fromstring(payload)
                if info.filename == "_rels/.rels":
                    etree.SubElement(root, f"{{{rel_ns}}}Relationship", Id="rIdCustom", Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties", Target="docProps/custom.xml")
                elif info.filename == "word/_rels/document.xml.rels":
                    etree.SubElement(root, f"{{{rel_ns}}}Relationship", Id="rIdComments", Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments", Target="comments.xml")
                elif info.filename == "[Content_Types].xml":
                    etree.SubElement(root, f"{{{content_ns}}}Override", PartName="/docProps/custom.xml", ContentType="application/vnd.openxmlformats-officedocument.custom-properties+xml")
                    etree.SubElement(root, f"{{{content_ns}}}Override", PartName="/word/comments.xml", ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml")
                else:
                    etree.SubElement(root, f"{{{app_ns}}}Company").text = "Private Nick"
                payload = etree.tostring(root, encoding="UTF-8", xml_declaration=True)
            authored.writestr(info, payload)
        authored.writestr("docProps/custom.xml", '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties"><property name="Owner"><value>Private Nick</value></property></Properties>')
        authored.writestr("word/comments.xml", '<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:comment w:id="0" w:author="Private Nick" w:initials="PN"><w:p><w:r><w:t>Keep this note</w:t></w:r></w:p></w:comment></w:comments>')

    cleaned = remove_all(source)
    report = analyze(cleaned)

    assert Document(cleaned).paragraphs[0].text == Document(source).paragraphs[0].text
    assert not {"Author", "Creator", "LastModifiedBy", "Company"} & {field.key for field in report.fields}
    with zipfile.ZipFile(cleaned) as archive:
        assert b"Private Nick" not in archive.read("docProps/app.xml")
        assert b"Private Nick" not in archive.read("docProps/custom.xml")
        comments = archive.read("word/comments.xml")
        assert b"Private Nick" not in comments
        assert b"Keep this note" in comments


def test_safe_extractall_rejects_path_traversal(tmp_path) -> None:
    archive_path = tmp_path / "malicious.zip"
    with zipfile.ZipFile(archive_path, "w") as archive:
        archive.writestr("../escape.txt", "bad")

    extract_root = tmp_path / "extract"
    extract_root.mkdir()
    with zipfile.ZipFile(archive_path) as archive:
        with pytest.raises(ValueError, match="Path traversal detected"):
            _safe_extractall(archive, extract_root)


def test_strip_docx_tracked_changes_keeps_non_tracked_tags(tmp_path) -> None:
    document_xml = tmp_path / "document.xml"
    document_xml.write_text(
        """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p>
      <w:ins><w:r><w:t>Inserted</w:t></w:r></w:ins>
      <w:instrText>DO_NOT_TOUCH</w:instrText>
      <w:del><w:r><w:delText>Deleted</w:delText></w:r></w:del>
    </w:p>
  </w:body>
</w:document>
""",
        encoding="utf-8",
    )

    _strip_docx_tracked_changes(document_xml)
    transformed = document_xml.read_text(encoding="utf-8")
    root = ET.fromstring(transformed)
    local_names = [element.tag.split("}")[-1] for element in root.iter()]

    assert "ins" not in local_names
    assert "del" not in local_names
    assert "instrText" in local_names
    assert "DO_NOT_TOUCH" in transformed
